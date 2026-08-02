# -*- coding: utf-8 -*-
"""Build a paper-results workbook from complete rerun records."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import json
from pathlib import Path
from typing import Callable, Dict, Final, List, Literal, Optional, Sequence, Tuple, Union

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from scripts.build_task_embedding_report import canonical_split_name, compute_metrics


SourceStatus = Literal["success", "failed"]
SplitName = Literal["Val", "Test1", "Test2", "Test3"]
JsonScalar = Union[None, bool, int, float, str]
JsonValue = Union[JsonScalar, List["JsonValue"], Dict[str, "JsonValue"]]

SOURCE_IDS: Final = (
    "BASE_TRANSFORMER", "BASE_CNN1D", "BASE_RNN", "BASE_LSTM", "BASE_GRU", "BASE_BILSTM",
    "TASK_BILSTM", "ORDINAL_BILSTM", "FULL_TRANSFORMER", "FULL_LSTM", "FULL_GRU",
    "FULL_CNN1D", "FULL_RNN", "FULL_BILSTM",
)
SPLITS: Final[Tuple[SplitName, ...]] = ("Val", "Test1", "Test2", "Test3")
OOD_SPLITS: Final[Tuple[SplitName, ...]] = ("Test1", "Test2", "Test3")
DEFAULT_TEMPLATE_PATH: Final = Path(__file__).resolve().parents[1] / "paper_best_results" / "论文实验最佳结果汇总.xlsx"


@dataclass(frozen=True)
class SourceRecord:
    source_id: str
    status: SourceStatus
    run_directory: Optional[Path]
    config_path: Optional[Path]
    test_results_path: Optional[Path]
    reason: Optional[str]


@dataclass(frozen=True)
class IncompleteRunRecordError(Exception):
    source_id: str
    detail: str

    def __str__(self) -> str:
        return f"{self.source_id}: {self.detail}"


@dataclass(frozen=True)
class RerunRecordError(Exception):
    detail: str

    def __str__(self) -> str:
        return self.detail


@dataclass(frozen=True)
class SplitPayload:
    labels: Tuple[int, ...]
    predictions: Tuple[int, ...]
    probabilities: Tuple[Tuple[float, ...], ...]


@dataclass(frozen=True)
class MetricValues:
    acc: float
    macro_f1: float
    g_mean: float
    auc_pr: float
    qwk: float


@dataclass(frozen=True)
class SourceMetrics:
    val: MetricValues
    test1: MetricValues
    test2: MetricValues
    test3: MetricValues

    def for_split(self, split: SplitName) -> MetricValues:
        if split == "Val":
            return self.val
        if split == "Test1":
            return self.test1
        if split == "Test2":
            return self.test2
        return self.test3


def build_results_workbook(
    records: Sequence[SourceRecord],
    output_path: Path,
    template_path: Path = DEFAULT_TEMPLATE_PATH,
) -> Path:
    """Copy the reference workbook and replace it with static rerun results."""
    if output_path.resolve() == template_path.resolve():
        raise RerunRecordError("output path must differ from template")
    records_by_id = _index_records(records)
    metrics_by_id = _collect_metrics(records_by_id)
    workbook = load_workbook(template_path, data_only=False)
    try:
        _write_result_sources(workbook["Result Sources"], records_by_id, metrics_by_id)
        _write_summaries(workbook, metrics_by_id)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    finally:
        workbook.close()
    return output_path


def _index_records(records: Sequence[SourceRecord]) -> Dict[str, SourceRecord]:
    records_by_id = {record.source_id: record for record in records}
    if len(records) != len(SOURCE_IDS) or set(records_by_id) != set(SOURCE_IDS):
        raise RerunRecordError("records must contain each of the 14 reference source IDs exactly once")
    return records_by_id


def _collect_metrics(records_by_id: Dict[str, SourceRecord]) -> Dict[str, Optional[SourceMetrics]]:
    metrics_by_id: Dict[str, Optional[SourceMetrics]] = {}
    for source_id in SOURCE_IDS:
        record = records_by_id[source_id]
        if record.status == "success":
            if record.test_results_path is None:
                raise IncompleteRunRecordError(source_id, "successful record has no test results path")
            metrics_by_id[source_id] = _load_source_metrics(record)
        else:
            if record.reason is None or not record.reason:
                raise RerunRecordError(f"{source_id} failed record requires a reason")
            metrics_by_id[source_id] = None
    return metrics_by_id


def _load_source_metrics(record: SourceRecord) -> SourceMetrics:
    payload = json.loads(record.test_results_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not all(isinstance(name, str) for name in payload):
        raise IncompleteRunRecordError(record.source_id, "test results must be a split-name object")
    splits: Dict[SplitName, SplitPayload] = {}
    for raw_name, raw_metrics in payload.items():
        split = _normalize_split(raw_name)
        if split is not None:
            if split in splits:
                raise IncompleteRunRecordError(record.source_id, f"duplicate {split} split")
            splits[split] = _parse_split_payload(raw_metrics, record.source_id, split)
    if set(splits) != set(SPLITS):
        raise IncompleteRunRecordError(record.source_id, "requires complete Val, Test1, Test2, and Test3 splits")
    return SourceMetrics(
        val=_compute(splits["Val"], record.source_id, "Val"),
        test1=_compute(splits["Test1"], record.source_id, "Test1"),
        test2=_compute(splits["Test2"], record.source_id, "Test2"),
        test3=_compute(splits["Test3"], record.source_id, "Test3"),
    )


def _normalize_split(raw_name: str) -> Optional[SplitName]:
    normalized = canonical_split_name(raw_name)
    if normalized == "InDist":
        return "Val"
    if normalized in OOD_SPLITS:
        return normalized
    return None


def _parse_split_payload(raw_metrics: JsonValue, source_id: str, split: SplitName) -> SplitPayload:
    if not isinstance(raw_metrics, dict):
        raise IncompleteRunRecordError(source_id, f"{split} split has no metric object")
    labels = _integer_values(raw_metrics.get("labels"), source_id, split, "labels")
    predictions = _integer_values(raw_metrics.get("predictions"), source_id, split, "predictions")
    probabilities = _probability_values(raw_metrics.get("probabilities"), source_id, split)
    if not labels or len(labels) != len(predictions) or len(labels) != len(probabilities):
        raise IncompleteRunRecordError(source_id, f"{split} labels, predictions, and probabilities must align")
    return SplitPayload(labels, predictions, probabilities)


def _integer_values(raw_values: Optional[JsonValue], source_id: str, split: SplitName, label: str) -> Tuple[int, ...]:
    if not isinstance(raw_values, list):
        raise IncompleteRunRecordError(source_id, f"{split} is missing {label}")
    values: List[int] = []
    for value in raw_values:
        if not isinstance(value, int) or isinstance(value, bool):
            raise IncompleteRunRecordError(source_id, f"{split} {label} must contain integers")
        values.append(value)
    return tuple(values)


def _probability_values(raw_values: Optional[JsonValue], source_id: str, split: SplitName) -> Tuple[Tuple[float, ...], ...]:
    if not isinstance(raw_values, list):
        raise IncompleteRunRecordError(source_id, f"{split} is missing probabilities")
    rows: List[Tuple[float, ...]] = []
    for raw_row in raw_values:
        if not isinstance(raw_row, list) or not raw_row:
            raise IncompleteRunRecordError(source_id, f"{split} probabilities must contain nonempty rows")
        row: List[float] = []
        for value in raw_row:
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise IncompleteRunRecordError(source_id, f"{split} probabilities must contain numbers")
            row.append(float(value))
        rows.append(tuple(row))
    return tuple(rows)


def _compute(payload: SplitPayload, source_id: str, split: SplitName) -> MetricValues:
    computed = compute_metrics(
        {"labels": list(payload.labels), "predictions": list(payload.predictions), "probabilities": [list(row) for row in payload.probabilities]},
        f"{source_id}/{split}",
    )
    return MetricValues(computed["ACC"], computed["F1 Macro"], computed["G-Mean"], computed["AUC-PR"], computed["QWK"])


def _write_result_sources(
    sheet: Worksheet,
    records_by_id: Dict[str, SourceRecord],
    metrics_by_id: Dict[str, Optional[SourceMetrics]],
) -> None:
    sheet.column_dimensions["P"].width = sheet.column_dimensions["O"].width
    sheet.auto_filter.ref = "A4:P60"
    sheet["P4"]._style = copy(sheet["O4"]._style)
    sheet["P4"].value = "Status"
    for source_index, source_id in enumerate(SOURCE_IDS):
        record = records_by_id[source_id]
        source_metrics = metrics_by_id[source_id]
        for split_index, split in enumerate(SPLITS):
            row = 5 + source_index * len(SPLITS) + split_index
            sheet.cell(row, 3).value = _path_value(record.run_directory)
            sheet.cell(row, 4).value = _path_value(record.test_results_path)
            sheet.cell(row, 5).value = _path_value(record.config_path)
            sheet.cell(row, 16)._style = copy(sheet.cell(row, 15)._style)
            sheet.cell(row, 16).value = _status_value(record)
            _write_metrics(sheet, row, 10, _source_split_values(source_metrics, split))


def _path_value(path: Optional[Path]) -> Optional[str]:
    return path.as_posix() if path is not None else None


def _status_value(record: SourceRecord) -> str:
    if record.status == "success":
        return "success"
    return f"failed: {record.reason}"


def _source_split_values(metrics: Optional[SourceMetrics], split: SplitName) -> Tuple[Optional[float], ...]:
    if metrics is None:
        return (None, None, None, None, None)
    values = metrics.for_split(split)
    return (values.acc, values.macro_f1, values.g_mean, values.auc_pr, values.qwk)


def _write_summaries(workbook, metrics_by_id: Dict[str, Optional[SourceMetrics]]) -> None:
    _write_distribution(workbook["Distribution Shift"], metrics_by_id["BASE_BILSTM"])
    _write_bilstm_configurations(workbook["BiLSTM Configurations"], metrics_by_id)
    _write_component_ablation(workbook["Component Ablation"], metrics_by_id)
    _write_encoder_comparison(workbook["Encoder Comparison"], metrics_by_id)


def _write_distribution(sheet: Worksheet, metrics: Optional[SourceMetrics]) -> None:
    for row, split in enumerate(SPLITS, start=5):
        _write_metrics(sheet, row, 2, _source_split_values(metrics, split))


def _write_bilstm_configurations(sheet: Worksheet, metrics_by_id: Dict[str, Optional[SourceMetrics]]) -> None:
    for row, source_id in enumerate(("BASE_BILSTM", "TASK_BILSTM", "ORDINAL_BILSTM", "FULL_BILSTM"), start=5):
        metrics = metrics_by_id[source_id]
        values = _split_f1_values(metrics) if metrics is not None else (None, None, None, None, None)
        _write_metrics(sheet, row, 2, values)


def _write_component_ablation(sheet: Worksheet, metrics_by_id: Dict[str, Optional[SourceMetrics]]) -> None:
    for row, source_id in enumerate(("BASE_BILSTM", "TASK_BILSTM", "ORDINAL_BILSTM", "FULL_BILSTM"), start=5):
        metrics = metrics_by_id[source_id]
        values = _component_values(metrics) if metrics is not None else (None, None, None, None, None, None)
        _write_metrics(sheet, row, 2, values)


def _write_encoder_comparison(sheet: Worksheet, metrics_by_id: Dict[str, Optional[SourceMetrics]]) -> None:
    pairs = (("BASE_TRANSFORMER", "FULL_TRANSFORMER"), ("BASE_CNN1D", "FULL_CNN1D"), ("BASE_RNN", "FULL_RNN"), ("BASE_LSTM", "FULL_LSTM"), ("BASE_GRU", "FULL_GRU"), ("BASE_BILSTM", "FULL_BILSTM"))
    for row, (base_id, full_id) in enumerate(pairs, start=5):
        base = metrics_by_id[base_id]
        full = metrics_by_id[full_id]
        values = _encoder_values(base, full) if base is not None and full is not None else (None, None, None, None, None, None)
        _write_metrics(sheet, row, 2, values)


def _split_f1_values(metrics: SourceMetrics) -> Tuple[float, float, float, float, float]:
    return tuple(metrics.for_split(split).macro_f1 for split in SPLITS) + (_ood_mean(metrics, lambda value: value.macro_f1),)


def _component_values(metrics: SourceMetrics) -> Tuple[float, float, float, float, float, float]:
    return (
        metrics.val.macro_f1,
        _ood_mean(metrics, lambda value: value.acc),
        _ood_mean(metrics, lambda value: value.macro_f1),
        _ood_mean(metrics, lambda value: value.g_mean),
        _ood_mean(metrics, lambda value: value.auc_pr),
        _ood_mean(metrics, lambda value: value.qwk),
    )


def _encoder_values(base: SourceMetrics, full: SourceMetrics) -> Tuple[float, float, float, float, float, float]:
    base_f1 = _ood_mean(base, lambda value: value.macro_f1)
    full_f1 = _ood_mean(full, lambda value: value.macro_f1)
    return (base_f1, full_f1, (full_f1 - base_f1) * 100, _ood_mean(full, lambda value: value.g_mean), _ood_mean(full, lambda value: value.auc_pr), _ood_mean(full, lambda value: value.qwk))


def _ood_mean(metrics: SourceMetrics, selector: Callable[[MetricValues], float]) -> float:
    return sum(selector(metrics.for_split(split)) for split in OOD_SPLITS) / len(OOD_SPLITS)


def _write_metrics(sheet: Worksheet, row: int, start_column: int, values: Sequence[Optional[float]]) -> None:
    for column, value in enumerate(values, start=start_column):
        sheet.cell(row, column).value = value
