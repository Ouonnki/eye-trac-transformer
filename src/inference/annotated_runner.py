# -*- coding: utf-8 -*-
"""Runner for annotated new-data validation and six-label inference."""

from __future__ import annotations

import csv
import json
import shutil
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

from openpyxl import Workbook

from src.inference.annotated_segments import (
    AnnotatedRecord,
    IssueRecord,
    build_annotated_records,
    expand_records_by_conditions,
)
from src.inference.new_data_model import (
    load_inference_model,
    predict_samples,
    resolve_device,
)
from src.inference.new_data_types import PredictionResult, TASK_SPECS


DEFAULT_DATA_DIR = Path("新数据+标注")
DEFAULT_OUTPUT_ROOT = Path("outputs/annotated_new_data_tests")
MANIFEST_CSV = "manifest.csv"
MANIFEST_XLSX = "manifest.xlsx"
PREDICTIONS_CSV = "predictions.csv"
PREDICTIONS_XLSX = "predictions.xlsx"
SUMMARY_SHEET = "summary"
ISSUES_SHEET = "issues"
PREDICTION_COLUMN = "预测结果"
BASE_COLUMNS = (
    "subject_id",
    "task_key",
    "task_id",
    "task_name",
    "condition_label",
    "condition_1",
    "condition_2",
    "condition_3",
    "condition_4",
    "condition_5",
    "segment_id",
    "segment_kind",
    "source_paths",
    "model_segment_count",
    "status",
    "warnings",
)
PREDICTION_COLUMNS = (
    "pred_class_0based",
    "pred_label",
    "prob_low",
    "prob_mid",
    "prob_high",
    PREDICTION_COLUMN,
)
ISSUE_COLUMNS = ("task_key", "subject_id", "segment_id", "source_paths", "reason", "warnings")


@dataclass(frozen=True)
class _OutputSpec:
    records: Sequence[AnnotatedRecord]
    issues: Sequence[IssueRecord]
    output_root: Path
    kind: str
    predictions: Sequence[PredictionResult] = ()


@dataclass(frozen=True)
class _WorkbookSpec:
    records: Sequence[AnnotatedRecord]
    rows: Sequence[Mapping]
    columns: Sequence[str]
    issues: Sequence[IssueRecord]


def run_annotated_validate(
    data_dir: Path = DEFAULT_DATA_DIR,
    output_root: Path = DEFAULT_OUTPUT_ROOT,
) -> Path:
    build = build_annotated_records(data_dir)
    records = expand_records_by_conditions(build.records, build.conditions_by_task)
    return _write_outputs(_OutputSpec(records, build.issues, output_root, "manifest"))


def run_annotated_inference(
    data_dir: Path,
    checkpoint_path: Path,
    *,
    output_root: Path = DEFAULT_OUTPUT_ROOT,
    device_name: str = "auto",
) -> Path:
    if checkpoint_path is None:
        raise ValueError("infer 模式必须提供 checkpoint")
    device = resolve_device(device_name)
    model, config = load_inference_model(checkpoint_path, device)
    screen_width, screen_height = _screen_size(config)
    build = build_annotated_records(data_dir, screen_width, screen_height)
    records = expand_records_by_conditions(build.records, build.conditions_by_task)
    predictions = predict_samples(model, config, _samples(records), device)
    spec = _OutputSpec(records, build.issues, output_root, "predictions", predictions)
    return _write_outputs(spec)


def _write_outputs(spec: _OutputSpec) -> Path:
    final_dir, temporary_dir = _create_temporary_output(spec.output_root)
    try:
        rows, columns = _rows_and_columns(spec.records, spec.predictions)
        _write_csv(temporary_dir / _csv_name(spec.kind), columns, rows)
        workbook = _WorkbookSpec(spec.records, rows, columns, spec.issues)
        _write_workbook(temporary_dir / _xlsx_name(spec.kind), workbook)
        temporary_dir.rename(final_dir)
    except Exception:
        shutil.rmtree(temporary_dir, ignore_errors=True)
        _remove_empty_directory(spec.output_root)
        raise
    return final_dir


def _rows_and_columns(
    records: Sequence[AnnotatedRecord],
    predictions: Sequence[PredictionResult],
) -> Tuple[Tuple[Dict, ...], Tuple[str, ...]]:
    if predictions:
        return _prediction_rows(records, predictions), BASE_COLUMNS + PREDICTION_COLUMNS
    return tuple(_base_row(record) for record in records), BASE_COLUMNS


def _prediction_rows(
    records: Sequence[AnnotatedRecord],
    predictions: Sequence[PredictionResult],
) -> Tuple[Dict, ...]:
    if len(records) != len(predictions):
        raise ValueError("预测结果数量与 manifest 行数不一致")
    rows = []
    for record, prediction in zip(records, predictions):
        row = _base_row(record)
        row.update(_prediction_columns(prediction))
        rows.append(row)
    return tuple(rows)


def _prediction_columns(prediction: PredictionResult) -> Dict:
    probabilities = prediction.probabilities
    return {
        "pred_class_0based": prediction.predicted_index,
        "pred_label": prediction.predicted_label,
        "prob_low": probabilities[0],
        "prob_mid": probabilities[1],
        "prob_high": probabilities[2],
        PREDICTION_COLUMN: prediction.output_value,
    }


def _base_row(record: AnnotatedRecord) -> Dict:
    sample = record.sample
    return {
        "subject_id": sample.subject_id,
        "task_key": sample.task_key,
        "task_id": sample.task_id,
        "task_name": sample.task_name,
        "condition_label": sample.condition_label,
        "condition_1": sample.task_conditions[0],
        "condition_2": sample.task_conditions[1],
        "condition_3": sample.task_conditions[2],
        "condition_4": sample.task_conditions[3],
        "condition_5": sample.task_conditions[4],
        "segment_id": record.segment_id,
        "segment_kind": record.segment_kind,
        "source_paths": _json_paths(record.source_paths),
        "model_segment_count": len(sample.segments),
        "status": record.status,
        "warnings": _json_tuple(record.warnings),
    }


def _write_workbook(path: Path, workbook_spec: _WorkbookSpec) -> None:
    workbook = Workbook()
    workbook.active.title = SUMMARY_SHEET
    _append_rows(workbook[SUMMARY_SHEET], workbook_spec.columns, workbook_spec.rows)
    for sheet_name, sheet_rows in _sheet_rows(workbook_spec.records, workbook_spec.rows).items():
        sheet = workbook.create_sheet(sheet_name)
        _append_rows(sheet, workbook_spec.columns, sheet_rows)
    issue_rows = _issue_rows(workbook_spec.issues)
    _append_rows(workbook.create_sheet(ISSUES_SHEET), ISSUE_COLUMNS, issue_rows)
    workbook.save(path)


def _sheet_rows(
    records: Sequence[AnnotatedRecord],
    rows: Sequence[Mapping],
) -> Dict[str, List[Mapping]]:
    grouped: Dict[str, List[Mapping]] = OrderedDict()
    for record, row in zip(records, rows):
        grouped.setdefault(_sheet_name(record), []).append(row)
    return grouped


def _sheet_name(record: AnnotatedRecord) -> str:
    spec = TASK_SPECS[record.sample.task_key]
    name = f"{spec.sheet_prefix}_{record.sample.condition_label}"
    if len(name) > 31:
        raise ValueError(f"Excel sheet 名称超过 31 字符: {name}")
    return name


def _append_rows(sheet, columns: Sequence[str], rows: Iterable[Mapping]) -> None:
    sheet.append(tuple(columns))
    for row in rows:
        sheet.append(tuple(row.get(column, "") for column in columns))


def _issue_rows(issues: Sequence[IssueRecord]) -> Tuple[Dict, ...]:
    return tuple(
        {
            "task_key": issue.task_key,
            "subject_id": issue.subject_id,
            "segment_id": issue.segment_id,
            "source_paths": _json_paths(issue.source_paths),
            "reason": issue.reason,
            "warnings": _json_tuple(issue.warnings),
        }
        for issue in issues
    )


def _create_temporary_output(output_root: Path) -> Tuple[Path, Path]:
    run_name = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    final_dir = output_root / run_name
    temporary_dir = output_root / f".{run_name}.tmp"
    if final_dir.exists() or temporary_dir.exists():
        raise FileExistsError(f"输出目录已存在: {final_dir}")
    output_root.mkdir(parents=True, exist_ok=True)
    temporary_dir.mkdir()
    return final_dir, temporary_dir


def _write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _screen_size(config: Mapping) -> Tuple[int, int]:
    sequence = config.get("sequence", {})
    try:
        return int(sequence["screen_width"]), int(sequence["screen_height"])
    except KeyError as error:
        raise ValueError("config.json 缺少 sequence.screen_width/screen_height") from error


def _csv_name(kind: str) -> str:
    if kind == "manifest":
        return MANIFEST_CSV
    return PREDICTIONS_CSV


def _xlsx_name(kind: str) -> str:
    if kind == "manifest":
        return MANIFEST_XLSX
    return PREDICTIONS_XLSX


def _samples(records: Sequence[AnnotatedRecord]) -> Tuple:
    return tuple(record.sample for record in records)


def _json_paths(paths: Sequence[Path]) -> str:
    return json.dumps(tuple(str(path) for path in paths), ensure_ascii=False)


def _json_tuple(values: Sequence[str]) -> str:
    return json.dumps(tuple(values), ensure_ascii=False)


def _remove_empty_directory(path: Path) -> None:
    if path.exists() and not any(path.iterdir()):
        path.rmdir()
