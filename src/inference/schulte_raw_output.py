# -*- coding: utf-8 -*-
"""CSV/XLSX output helpers for Schulte raw validation and inference."""

from __future__ import annotations

import csv
import json
import shutil
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

from openpyxl import Workbook

from src.inference.schulte_raw_loader import READY_STATUS
from src.inference.schulte_raw_types import SchulteBuild, SchulteIssue, SchulteRecord


MANIFEST_CSV = "manifest.csv"
MANIFEST_XLSX = "manifest.xlsx"
PREDICTIONS_CSV = "predictions.csv"
PREDICTIONS_XLSX = "predictions.xlsx"
SUMMARY_SHEET = "summary"
ISSUES_SHEET = "issues"
MANIFEST_SHEET = "manifest"
PREDICTIONS_SHEET = "predictions"
PREDICTION_COLUMN = "预测结果"
BASE_COLUMNS = (
    "subject_id",
    "task_id",
    "source_path",
    "grid_size",
    "number_range",
    "condition_1",
    "condition_2",
    "condition_3",
    "condition_4",
    "condition_5",
    "raw_duration",
    "reported_error_count",
    "actual_error_count",
    "click_count",
    "segment_count",
    "gaze_count",
    "timestamp_fix_count",
    "clipped_coordinate_count",
    "status",
    "warnings",
)
PREDICTION_COLUMNS = (
    "pred_class_0based",
    "pred_label",
    "prob_low",
    "prob_mid",
    "prob_high",
    PREDICTION_COLUMN,
)
ISSUE_COLUMNS = ("subject_id", "task_id", "source_path", "reason", "warnings")


@dataclass(frozen=True)
class SchulteOutputRequest:
    build: SchulteBuild
    output_root: Path
    kind: str
    screen_size: Tuple[int, int]
    predictions: Sequence = ()


def write_schulte_outputs(request: SchulteOutputRequest) -> Path:
    final_dir, temporary_dir = _create_temporary_output(request.output_root)
    try:
        rows, columns = _rows_and_columns(request)
        _write_csv(temporary_dir / _csv_name(request.kind), columns, rows)
        _write_workbook(
            temporary_dir / _xlsx_name(request.kind),
            request,
            rows,
            columns=columns,
        )
        temporary_dir.rename(final_dir)
    except Exception:
        shutil.rmtree(temporary_dir, ignore_errors=True)
        _remove_empty_directory(request.output_root)
        raise
    return final_dir


def _rows_and_columns(request: SchulteOutputRequest) -> Tuple[Tuple[Dict, ...], Tuple[str, ...]]:
    if request.predictions:
        rows = _prediction_rows(request.build.records, request.predictions)
        return rows, BASE_COLUMNS + PREDICTION_COLUMNS
    return tuple(_base_row(record) for record in request.build.records), BASE_COLUMNS


def _prediction_rows(records: Sequence[SchulteRecord], predictions: Sequence) -> Tuple[Dict, ...]:
    ready_records = tuple(record for record in records if record.status == READY_STATUS)
    if len(ready_records) != len(predictions):
        raise ValueError("预测结果数量与 ready 记录数不一致")
    rows = []
    for record, prediction in zip(ready_records, predictions):
        row = _base_row(record)
        row.update(_prediction_columns(prediction))
        rows.append(row)
    return tuple(rows)


def _base_row(record: SchulteRecord) -> Dict:
    question = record.question
    return {
        "subject_id": record.subject_id,
        "task_id": record.task_id,
        "source_path": str(record.source_path),
        "grid_size": question.grid_size,
        "number_range": question.number_range_text,
        "condition_1": question.task_conditions[0],
        "condition_2": question.task_conditions[1],
        "condition_3": question.task_conditions[2],
        "condition_4": question.task_conditions[3],
        "condition_5": question.task_conditions[4],
        "raw_duration": record.raw_duration,
        "reported_error_count": record.reported_error_count,
        "actual_error_count": record.actual_error_count,
        "click_count": record.click_count,
        "segment_count": record.segment_count,
        "gaze_count": record.gaze_count,
        "timestamp_fix_count": record.timestamp_fix_count,
        "clipped_coordinate_count": record.clipped_coordinate_count,
        "status": record.status,
        "warnings": _json_tuple(record.warnings),
    }


def _prediction_columns(prediction) -> Dict:
    probabilities = prediction.probabilities
    return {
        "pred_class_0based": prediction.predicted_index,
        "pred_label": prediction.predicted_label,
        "prob_low": probabilities[0],
        "prob_mid": probabilities[1],
        "prob_high": probabilities[2],
        PREDICTION_COLUMN: prediction.output_value,
    }


def _write_workbook(
    path: Path,
    request: SchulteOutputRequest,
    rows: Sequence[Mapping],
    *,
    columns: Sequence[str],
) -> None:
    workbook = Workbook()
    workbook.active.title = SUMMARY_SHEET
    _append_rows(workbook[SUMMARY_SHEET], ("metric", "value"), _summary_rows(request))
    main_sheet = PREDICTIONS_SHEET if request.kind == "predictions" else MANIFEST_SHEET
    _append_rows(workbook.create_sheet(main_sheet), columns, rows)
    _write_question_sheets(workbook, rows, columns)
    _append_rows(workbook.create_sheet(ISSUES_SHEET), ISSUE_COLUMNS, _issue_rows(request.build.issues))
    workbook.save(path)


def _write_question_sheets(
    workbook: Workbook,
    rows: Sequence[Mapping],
    columns: Sequence[str],
) -> None:
    for sheet_name, sheet_rows in _question_sheet_rows(rows).items():
        _append_rows(workbook.create_sheet(sheet_name), columns, sheet_rows)


def _summary_rows(request: SchulteOutputRequest) -> Tuple[Dict, ...]:
    records = request.build.records
    return (
        {"metric": "record_count", "value": len(records)},
        {"metric": "ready_count", "value": _ready_count(records)},
        {"metric": "issue_count", "value": len(request.build.issues)},
        {"metric": "warning_record_count", "value": _warning_count(records)},
        {"metric": "timestamp_fix_count", "value": sum(record.timestamp_fix_count for record in records)},
        {"metric": "clipped_coordinate_count", "value": sum(record.clipped_coordinate_count for record in records)},
        {"metric": "screen_width", "value": request.screen_size[0]},
        {"metric": "screen_height", "value": request.screen_size[1]},
    )


def _question_sheet_rows(rows: Sequence[Mapping]) -> Mapping[str, List[Mapping]]:
    grouped: Dict[str, List[Mapping]] = OrderedDict()
    for row in rows:
        grouped.setdefault(f"q{int(row['task_id']):02d}", []).append(row)
    return grouped


def _issue_rows(issues: Sequence[SchulteIssue]) -> Tuple[Dict, ...]:
    return tuple(
        {
            "subject_id": issue.subject_id,
            "task_id": issue.task_id,
            "source_path": str(issue.source_path),
            "reason": issue.reason,
            "warnings": _json_tuple(issue.warnings),
        }
        for issue in issues
    )


def _append_rows(sheet, columns: Sequence[str], rows: Iterable[Mapping]) -> None:
    sheet.append(tuple(columns))
    for row in rows:
        sheet.append(tuple(row.get(column, "") for column in columns))


def _write_csv(path: Path, columns: Sequence[str], rows: Sequence[Mapping]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _create_temporary_output(output_root: Path) -> Tuple[Path, Path]:
    run_name = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    final_dir = output_root / run_name
    temporary_dir = output_root / f".{run_name}.tmp"
    if final_dir.exists() or temporary_dir.exists():
        raise FileExistsError(f"输出目录已存在: {final_dir}")
    output_root.mkdir(parents=True, exist_ok=True)
    temporary_dir.mkdir()
    return final_dir, temporary_dir


def _ready_count(records: Sequence[SchulteRecord]) -> int:
    return sum(record.status == READY_STATUS for record in records)


def _warning_count(records: Sequence[SchulteRecord]) -> int:
    return sum(bool(record.warnings) for record in records)


def _csv_name(kind: str) -> str:
    return PREDICTIONS_CSV if kind == "predictions" else MANIFEST_CSV


def _xlsx_name(kind: str) -> str:
    return PREDICTIONS_XLSX if kind == "predictions" else MANIFEST_XLSX


def _json_tuple(values: Sequence[str]) -> str:
    return json.dumps(tuple(values), ensure_ascii=False)


def _remove_empty_directory(path: Path) -> None:
    if path.exists() and not any(path.iterdir()):
        path.rmdir()
