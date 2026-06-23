# -*- coding: utf-8 -*-
"""Tests for annotated new-data six-label validation and inference."""

import csv
import json
import tempfile
import unittest
from pathlib import Path

import torch
from openpyxl import load_workbook

from src.inference.annotated_runner import (
    run_annotated_inference,
    run_annotated_validate,
)
from src.inference.annotated_segments import (
    build_situation_awareness_records,
    build_spot_subject_records,
    parse_condition_combinations,
)
from src.models.task_level_model import TaskLevelEncoder


SA_DIR = "新数据+标注-情景意识眼动新数据（120Hz，降采样一半使用）（3_4_5_6，1_0，1，0，1）"
SPOT_DIR = "新数据+标注-找不同数据（眼动+行为_按题目）参数-（3_4_5，1，0_1，1，0）"
SPOT_EYE = "找不同眼动（120Hz，降采样一半使用）（关键事件：TaskStart和ROIHitClick）"
SPOT_BEHAVIOR = "找不同行为数据（按题目）"
CSV_FIELDS = [
    "Recording timestamp",
    "Event",
    "Event value",
    "Gaze point X",
    "Gaze point Y",
]


class AnnotatedNewDataRunnerTest(unittest.TestCase):
    def test_condition_combinations_parse_expected_counts(self):
        sa = parse_condition_combinations(SA_DIR)
        spot = parse_condition_combinations(SPOT_DIR)

        self.assertEqual(len(sa), 8)
        self.assertIn((3, 1, 1, 0, 1), sa)
        self.assertIn((6, 0, 1, 0, 1), sa)
        self.assertEqual(len(spot), 6)
        self.assertIn((3, 1, 0, 1, 0), spot)
        self.assertIn((5, 1, 1, 1, 0), spot)

    def test_situation_awareness_builds_five_segments_and_combined_sample(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "0429001_找宝石.csv"
            write_sa_csv(path)

            records = build_situation_awareness_records(path, 1920, 1080)

        self.assertEqual([record.segment_id for record in records], [
            "L1",
            "L2",
            "L3",
            "L4",
            "L5",
            "combined",
        ])
        self.assertEqual([len(record.sample.segments) for record in records], [
            1,
            1,
            1,
            1,
            1,
            5,
        ])
        self.assertEqual(records[-1].segment_kind, "combined")

    def test_spot_builds_topic_segments_and_warns_when_roi_hit_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            paths = {}
            for index in range(1, 6):
                path = root / f"T{index:03d}" / f"0429001_{index:03d}.csv"
                write_spot_csv(path, include_roi=index != 2)
                paths[f"T{index:03d}"] = path

            records = build_spot_subject_records("0429001", paths, 1920, 1080)

        self.assertEqual([record.segment_id for record in records], [
            "T001",
            "T002",
            "T003",
            "T004",
            "T005",
            "combined",
        ])
        self.assertEqual(len(records[0].sample.segments), 2)
        self.assertEqual(len(records[1].sample.segments), 1)
        self.assertIn("ROIHitClick", records[1].warnings[0])
        self.assertEqual(len(records[-1].sample.segments), 9)

    def test_spot_warns_and_skips_too_short_internal_segment(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            paths = {}
            for index in range(1, 6):
                path = root / f"T{index:03d}" / f"0429001_{index:03d}.csv"
                write_spot_csv(path, include_roi=True, short_first_interval=index == 1)
                paths[f"T{index:03d}"] = path

            records = build_spot_subject_records("0429001", paths, 1920, 1080)

        self.assertEqual(len(records[0].sample.segments), 1)
        self.assertTrue(any("跳过" in warning for warning in records[0].warnings))

    def test_validate_writes_csv_xlsx_condition_sheets_and_issues_sheet(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = build_dataset(Path(temp_dir) / "新数据+标注")
            output = run_annotated_validate(
                data_dir=data_dir,
                output_root=Path(temp_dir) / "outputs",
            )

            csv_path = output / "manifest.csv"
            csv_exists = csv_path.is_file()
            row_count = count_csv_rows(csv_path)
            workbook = load_workbook(output / "manifest.xlsx", read_only=True)
            sheetnames = tuple(workbook.sheetnames)
            workbook.close()

        self.assertTrue(csv_exists)
        self.assertIn("summary", sheetnames)
        self.assertIn("issues", sheetnames)
        self.assertIn("sa_3-1-1-0-1", sheetnames)
        self.assertIn("spot_5-1-1-1-0", sheetnames)
        self.assertEqual(row_count, 84)

    def test_infer_writes_predictions_with_six_label_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir = build_dataset(root / "新数据+标注")
            checkpoint = root / "model" / "best_model.pt"
            write_checkpoint(checkpoint)

            output = run_annotated_inference(
                data_dir=data_dir,
                checkpoint_path=checkpoint,
                output_root=root / "outputs",
                device_name="cpu",
            )
            csv_path = output / "predictions.csv"
            csv_exists = csv_path.is_file()
            header = read_csv_header(csv_path)
            row_count = count_csv_rows(csv_path)
            workbook = load_workbook(output / "predictions.xlsx", read_only=True)
            sheetnames = tuple(workbook.sheetnames)
            workbook.close()

        self.assertTrue(csv_exists)
        self.assertIn("预测结果", header)
        self.assertIn("sa_6-0-1-0-1", sheetnames)
        self.assertEqual(row_count, 84)


def write_sa_csv(path: Path) -> None:
    fields = CSV_FIELDS + ["behavior_route_id"]
    rows = []
    timestamp = 0
    for level in range(1, 6):
        for offset in range(4):
            timestamp += 16_667
            rows.append([timestamp, "", "", 100 + offset, 200, f"L{level}_R1"])
    write_csv(path, fields, rows)


def write_spot_csv(path: Path, include_roi: bool, short_first_interval: bool = False) -> None:
    if short_first_interval:
        return write_short_spot_csv(path)
    rows = [["0", "TaskStart1", "", "", ""]]
    rows += [[str(value), "", "", "100", "200"] for value in (10_000, 20_000)]
    if include_roi:
        rows.append(["30_000", "ROIHitClick", "1", "", ""])
        rows += [[str(value), "", "", "120", "220"] for value in (40_000, 50_000)]
    rows += [[str(value), "", "", "140", "240"] for value in (60_000, 70_000)]
    rows.append(["80_000", "TaskEnd1", "", "", ""])
    write_csv(path, CSV_FIELDS, rows)


def write_short_spot_csv(path: Path) -> None:
    rows = [
        ["0", "TaskStart1", "", "", ""],
        ["1_000", "", "", "100", "200"],
        ["2_000", "ROIHitClick", "1", "", ""],
        ["20_000", "", "", "120", "220"],
        ["40_000", "", "", "130", "230"],
        ["60_000", "", "", "140", "240"],
        ["80_000", "TaskEnd1", "", "", ""],
    ]
    write_csv(path, CSV_FIELDS, rows)


def write_csv(path: Path, fields, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(fields)
        writer.writerows(rows)


def build_dataset(data_dir: Path) -> Path:
    sa_dir = data_dir / SA_DIR
    write_sa_csv(sa_dir / "0429001_找宝石.csv")
    write_summary(sa_dir / "批量标记_summary.csv")
    eye_root = data_dir / SPOT_DIR / SPOT_EYE
    behavior_root = data_dir / SPOT_DIR / SPOT_BEHAVIOR
    for index in range(1, 6):
        write_spot_csv(
            eye_root / f"T{index:03d}" / f"0429001_{index:03d}.csv",
            include_roi=index != 2,
        )
        write_dummy_behavior(behavior_root / f"T{index:03d}" / f"0429001_{index:03d}_behavior.xlsx")
    return data_dir


def write_summary(path: Path) -> None:
    fields = ["participant_id", "eye_file", "status", "warnings"]
    rows = [["0429001", "0429001_找宝石.csv", "完成", ""]]
    write_csv(path, fields, rows)


def write_dummy_behavior(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"placeholder")


def count_csv_rows(path: Path) -> int:
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return sum(1 for _ in csv.DictReader(file_obj))


def read_csv_header(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return next(csv.reader(file_obj))


def write_checkpoint(checkpoint: Path) -> None:
    checkpoint.parent.mkdir(parents=True)
    config = model_config()
    (checkpoint.parent / "config.json").write_text(
        json.dumps(config),
        encoding="utf-8",
    )
    model = TaskLevelEncoder(**model_options(config))
    torch.save({"model_state_dict": model.state_dict()}, checkpoint)


def model_config():
    return {
        "sequence": {"max_seq_len": 8, "max_segments": 25, "screen_width": 1920, "screen_height": 1080, "input_dim": 7},
        "model": {
            "segment_d_model": 4,
            "segment_nhead": 1,
            "segment_num_layers": 1,
            "segment_encoder_type": "bilstm",
            "segment_rnn_hidden_size": 4,
            "attention_dim": 4,
            "task_embedding_dim": 1,
            "use_task_embedding": True,
            "dropout": 0.0,
            "num_classes": 3,
            "head_type": "ordinal",
        },
        "training": {"batch_size": 2, "ordinal_thresholds": [0.47, 0.5]},
    }


def model_options(config):
    sequence = config["sequence"]
    model = config["model"]
    return {
        "input_dim": sequence["input_dim"],
        "max_seq_len": sequence["max_seq_len"],
        "max_segments": sequence["max_segments"],
        "segment_d_model": model["segment_d_model"],
        "segment_nhead": model["segment_nhead"],
        "segment_num_layers": model["segment_num_layers"],
        "segment_encoder_type": model["segment_encoder_type"],
        "segment_rnn_hidden_size": model["segment_rnn_hidden_size"],
        "attention_dim": model["attention_dim"],
        "task_embedding_dim": model["task_embedding_dim"],
        "use_task_embedding": model["use_task_embedding"],
        "dropout": model["dropout"],
        "num_classes": model["num_classes"],
        "head_type": model["head_type"],
    }


if __name__ == "__main__":
    unittest.main()
