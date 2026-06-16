# -*- coding: utf-8 -*-
"""End-to-end tests for new-data inference output."""

import csv
import json
import tempfile
import unittest
from pathlib import Path

import torch
from openpyxl import load_workbook

from src.inference.new_data_dataset import NewDataInferenceDataset, inference_collate_fn
from src.inference.new_data_runner import (
    DEFAULT_DATA_DIR,
    DEFAULT_CHECKPOINT_PATH,
    discover_task_files,
    expand_samples_by_conditions,
    load_inference_model,
    parse_task_conditions,
    run_new_data_inference,
)
from src.inference.new_data_model import decode_ordinal_predictions
from src.inference.new_data_types import (
    InferenceSample,
    PredictionResult,
    TASK_SPECS,
)
from src.models.task_level_model import TaskLevelEncoder


CSV_FIELDS = [
    "Recording timestamp",
    "Event",
    "Event value",
    "Gaze point X",
    "Gaze point Y",
]


class NewDataRunnerTest(unittest.TestCase):
    def test_unlabeled_dataset_and_collate(self):
        sample = self._sample()
        dataset = NewDataInferenceDataset((sample,), max_seq_len=4, input_dim=7)

        batch = inference_collate_fn([dataset[0]])

        self.assertNotIn("labels", batch)
        self.assertEqual(tuple(batch["segments"].shape), (1, 1, 4, 7))
        self.assertEqual(batch["subject_ids"], ["0429324"])
        self.assertEqual(batch["task_names"], ["复杂问题解决任务"])

    def test_missing_checkpoint_and_config_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            checkpoint = Path(temp_dir) / "best_model.pt"
            with self.assertRaises(FileNotFoundError):
                load_inference_model(checkpoint, torch.device("cpu"))
            checkpoint.touch()
            with self.assertRaisesRegex(FileNotFoundError, "config.json"):
                load_inference_model(checkpoint, torch.device("cpu"))

    def test_incompatible_checkpoint_fails(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            checkpoint = Path(temp_dir) / "best_model.pt"
            self._write_config(checkpoint.parent / "config.json")
            torch.save({"model_state_dict": {"wrong.weight": torch.ones(1)}}, checkpoint)
            with self.assertRaises(RuntimeError):
                load_inference_model(checkpoint, torch.device("cpu"))

    def test_ordinal_decoding_maps_internal_classes_to_one_two_three(self):
        logits = torch.tensor([[-10.0, -10.0], [10.0, -10.0], [10.0, 10.0]])
        config = {"training": {"ordinal_thresholds": [0.47, 0.5]}}

        predictions, probabilities = decode_ordinal_predictions(logits, config)
        output_values = [
            PredictionResult(self._sample(), value, ("低", "中", "高")[value], probs)
            .output_value
            for value, probs in zip(predictions, probabilities)
        ]

        self.assertEqual(predictions, (0, 1, 2))
        self.assertEqual(output_values, [1, 2, 3])
        for probability_row in probabilities:
            self.assertAlmostEqual(sum(probability_row), 1.0, places=6)

    def test_parse_task_conditions_from_directory_name(self):
        conditions = parse_task_conditions("找不同任务（3_4_5，1，0_1，1，0）")

        self.assertEqual(len(conditions), 6)
        self.assertIn((3, 1, 0, 1, 0), conditions)
        self.assertIn((5, 1, 1, 1, 0), conditions)

    def test_parse_task_conditions_accepts_colon_separator(self):
        conditions = parse_task_conditions("找不同任务（3:4:5，1，0:1，1，0）")

        self.assertEqual(len(conditions), 6)
        self.assertIn((3, 1, 0, 1, 0), conditions)
        self.assertIn((5, 1, 1, 1, 0), conditions)

    def test_discover_task_files_excludes_schulte_and_keeps_condition_counts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / "新数据（40人）"
            self._write_many_task_files(data_dir, subject_count=41)
            (data_dir / "舒尔特方格任务" / "舒尔特方格眼动数据").mkdir(parents=True)

            task_files = discover_task_files(data_dir)

        condition_counts = {
            task_set.spec.key: len(task_set.conditions)
            for task_set in task_files
        }
        file_counts = {
            task_set.spec.key: len(task_set.paths)
            for task_set in task_files
        }
        self.assertEqual(condition_counts, {
            "complex": 4,
            "situation_awareness": 4,
            "spot_difference": 6,
        })
        self.assertEqual(file_counts, {
            "complex": 41,
            "situation_awareness": 41,
            "spot_difference": 41,
        })

    def test_discover_task_files_requires_complete_subject_sets(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir) / "新数据（40人）"
            sources = self._write_many_task_files(data_dir, subject_count=2)
            next(path for path in sources if "找不同任务" in str(path)).unlink()

            with self.assertRaisesRegex(ValueError, "被试集合不完整"):
                discover_task_files(data_dir)

    def test_expand_samples_by_conditions_reuses_source_segments(self):
        sample = self._sample()
        expanded = expand_samples_by_conditions(
            (sample,),
            {
                "complex": ((1, 0, 0, 0, 0), (2, 0, 0, 0, 0)),
            },
        )

        self.assertEqual(len(expanded), 2)
        self.assertEqual(expanded[0].segments, sample.segments)
        self.assertEqual(expanded[0].task_conditions, (1, 0, 0, 0, 0))
        self.assertEqual(expanded[1].task_conditions, (2, 0, 0, 0, 0))

    def test_end_to_end_writes_single_workbook_with_condition_sheets(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir = root / "新数据（40人）"
            self._write_many_task_files(data_dir, subject_count=2)
            checkpoint = root / "model" / "best_model.pt"
            self._write_checkpoint(checkpoint)

            output_dir = run_new_data_inference(
                data_dir=data_dir,
                checkpoint_path=checkpoint,
                output_root=root / "outputs",
                device_name="cpu",
            )

            workbook_path = output_dir / "predictions.xlsx"
            self.assertTrue(workbook_path.is_file())
            self.assertFalse(any(output_dir.rglob("*_task.csv")))
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            self.assertEqual(len(workbook.sheetnames), 15)
            self.assertIn("summary", workbook.sheetnames)
            self.assertIn("complex_1-0-0-0-0", workbook.sheetnames)
            self.assertIn("spot_5-1-1-1-0", workbook.sheetnames)
            self.assertEqual(workbook["summary"].max_row, 29)
            self.assertEqual(workbook["complex_1-0-0-0-0"].max_row, 3)

    def test_failed_file_does_not_publish_output_directory(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir = root / "新数据（40人）"
            sources = self._write_many_task_files(data_dir, subject_count=1)
            self._write_csv(sources[0], [["0", "TaskStart1", "", "", ""]])
            checkpoint = root / "model" / "best_model.pt"
            self._write_checkpoint(checkpoint)
            output_root = root / "outputs"

            with self.assertRaises(ValueError):
                run_new_data_inference(data_dir, checkpoint, output_root, "cpu")

            self.assertFalse(output_root.exists())

    def test_default_checkpoint_matches_required_path(self):
        expected = (
            "outputs/task_level_full_data/"
            "task_level_ordinal_manual11_bilstm_full_data_20260607_233158/"
            "best_model.pt"
        )
        self.assertEqual(str(DEFAULT_CHECKPOINT_PATH), expected)
        self.assertEqual(str(DEFAULT_DATA_DIR), "新数据（40人）")

    def _sample(self):
        features = torch.zeros((2, 7), dtype=torch.float32).numpy()
        return InferenceSample(
            subject_id="0429324",
            task_id=101,
            task_name="复杂问题解决任务",
            task_conditions=(4, 0, 0, 0, 0),
            segments=(features,),
            source_path=Path("sample.csv"),
            task_key="complex",
        )

    def _write_many_task_files(self, data_dir, subject_count):
        sources = []
        subject_ids = [f"04293{index:02d}" for index in range(subject_count)]
        for spec in TASK_SPECS.values():
            eye_dir = data_dir / spec.directory_name / "眼动数据"
            eye_dir.mkdir(parents=True)
            for subject_id in subject_ids:
                path = eye_dir / f"{subject_id}_task.csv"
                self._write_task_csv(path, spec.key)
                sources.append(path)
        return sources

    def _write_task_csv(self, path, task_key):
        gaze = [[str(value), "", "", "100", "200"] for value in range(1_000, 90_000, 8_333)]
        if task_key == "complex":
            rows = [["0", "TaskStart1", "", "", ""]] + gaze
            rows += [["80_000", "MouseEvent", "Down, Left", "", ""]]
        elif task_key == "situation_awareness":
            rows = [["0", "TaskStart1", "", "", ""]] + gaze
            rows += [["70_000", "KeyboardEvent", "Enter", "", ""]]
        else:
            rows = [["0", "TaskStart1", "", "", ""]] + gaze
            rows += [
                ["70_000", "MouseEvent", "Down, Left", "", ""],
                ["80_000", "TaskEnd1", "", "", ""],
            ]
        self._write_csv(path, rows)

    def _write_csv(self, path, rows):
        with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
            writer = csv.writer(file_obj)
            writer.writerow(CSV_FIELDS)
            writer.writerows(rows)

    def _write_checkpoint(self, checkpoint):
        checkpoint.parent.mkdir(parents=True)
        config = self._write_config(checkpoint.parent / "config.json")
        model_config = config["model"]
        sequence = config["sequence"]
        model = TaskLevelEncoder(
            input_dim=sequence["input_dim"],
            max_seq_len=sequence["max_seq_len"],
            max_segments=sequence["max_segments"],
            segment_d_model=model_config["segment_d_model"],
            segment_nhead=model_config["segment_nhead"],
            segment_num_layers=model_config["segment_num_layers"],
            segment_encoder_type=model_config["segment_encoder_type"],
            segment_rnn_hidden_size=model_config["segment_rnn_hidden_size"],
            attention_dim=model_config["attention_dim"],
            task_embedding_dim=model_config["task_embedding_dim"],
            use_task_embedding=model_config["use_task_embedding"],
            dropout=model_config["dropout"],
            num_classes=model_config["num_classes"],
            head_type=model_config["head_type"],
        )
        torch.save({"model_state_dict": model.state_dict()}, checkpoint)

    def _write_config(self, path):
        config = {
            "sequence": {
                "max_seq_len": 8,
                "max_segments": 5,
                "screen_width": 1920,
                "screen_height": 1080,
                "input_dim": 7,
            },
            "model": {
                "segment_d_model": 4,
                "segment_nhead": 1,
                "segment_num_layers": 1,
                "segment_encoder_type": "bilstm",
                "segment_rnn_hidden_size": 4,
                "attention_dim": 4,
                "task_embedding_dim": 1,
                "use_task_embedding": True,
                "dropout": 0.0,
                "num_classes": 3,
                "head_type": "ordinal",
            },
            "training": {"batch_size": 2, "ordinal_thresholds": [0.47, 0.5]},
        }
        path.write_text(json.dumps(config), encoding="utf-8")
        return config


if __name__ == "__main__":
    unittest.main()
