# -*- coding: utf-8 -*-
"""Regression coverage for rerun paper-results workbooks."""

from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

from openpyxl import load_workbook

from paper_experiment_rerun.build_results_workbook import (
    IncompleteRunRecordError,
    RerunRecordError,
    SourceRecord,
    build_results_workbook,
)
from scripts.build_task_embedding_report import compute_metrics


SOURCE_IDS = (
    "BASE_TRANSFORMER",
    "BASE_CNN1D",
    "BASE_RNN",
    "BASE_LSTM",
    "BASE_GRU",
    "BASE_BILSTM",
    "TASK_BILSTM",
    "ORDINAL_BILSTM",
    "FULL_TRANSFORMER",
    "FULL_LSTM",
    "FULL_GRU",
    "FULL_CNN1D",
    "FULL_RNN",
    "FULL_BILSTM",
)
SPLITS = ("Val", "Test1", "Test2", "Test3")
REFERENCE_WORKBOOK = (
    Path(__file__).resolve().parents[1]
    / "paper_best_results"
    / "论文实验最佳结果汇总.xlsx"
)


class PaperExperimentReportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.payloads: Dict[str, Dict[str, Dict[str, Union[List[int], List[List[float]]]]]] = {}
        self.addCleanup(self.temp_dir.cleanup)

    def test_builds_static_rerun_workbook_and_marks_failed_source(self) -> None:
        # Given
        records = self._records(failed_source_id="FULL_CNN1D")
        output_path = self.root / "reports" / "batch-001" / "论文实验最佳结果汇总.xlsx"
        template_hash = hashlib.sha256(REFERENCE_WORKBOOK.read_bytes()).hexdigest()

        # When
        build_results_workbook(records, output_path, REFERENCE_WORKBOOK)

        # Then
        self.assertEqual(template_hash, hashlib.sha256(REFERENCE_WORKBOOK.read_bytes()).hexdigest())
        workbook = load_workbook(output_path, data_only=False)
        reference = load_workbook(REFERENCE_WORKBOOK, data_only=False)
        self.addCleanup(workbook.close)
        self.addCleanup(reference.close)
        self.assertEqual(workbook.sheetnames, reference.sheetnames)
        for sheet_name in workbook.sheetnames:
            self.assertEqual(workbook[sheet_name].freeze_panes, "A5")
            self.assertFalse(workbook[sheet_name].sheet_view.showGridLines)
            self.assertEqual(workbook[sheet_name].page_setup.orientation, "landscape")
            self.assertEqual(workbook[sheet_name].page_setup.fitToWidth, 1)
            self.assertEqual(workbook[sheet_name].page_setup.fitToHeight, 0)
        sources = workbook["Result Sources"]
        self.assertEqual(sources.max_row - 4, 56)
        self.assertEqual(sources.max_column, 16)
        self.assertEqual(sources["P4"].value, "Status")
        self.assertEqual(sources.auto_filter.ref, "A4:P60")
        self.assertEqual({str(item) for item in sources.merged_cells.ranges}, {"A1:O1", "A2:O2"})
        self.assertEqual(sources["P4"].style_id, sources["O4"].style_id)
        self.assertEqual(sources["P5"].style_id, sources["O5"].style_id)
        self.assertEqual(sources["C25"].value, (self.root / "BASE_BILSTM").as_posix())
        self.assertEqual(sources["D25"].value, (self.root / "BASE_BILSTM" / "test_results.json").as_posix())
        self.assertEqual(sources["E25"].value, (self.root / "BASE_BILSTM" / "config.json").as_posix())
        self.assertEqual(sources["P25"].value, "success")
        for row in range(49, 53):
            self.assertEqual(sources.cell(row, 3).value, (self.root / "FULL_CNN1D").as_posix())
            self.assertEqual(sources.cell(row, 4).value, (self.root / "FULL_CNN1D" / "test_results.json").as_posix())
            self.assertEqual(sources.cell(row, 5).value, (self.root / "FULL_CNN1D" / "config.json").as_posix())
            self.assertEqual([sources.cell(row, column).value for column in range(10, 15)], [None] * 5)
            self.assertEqual(sources.cell(row, 16).value, "failed: GPU memory exhausted")
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row))
        self._assert_summary_values(workbook)

    def test_successful_record_requires_every_complete_split(self) -> None:
        # Given
        records = self._records()
        incomplete_path = self.root / "TASK_BILSTM" / "test_results.json"
        incomplete_payload = self.payloads["TASK_BILSTM"].copy()
        incomplete_payload.pop("Test3")
        incomplete_path.write_text(json.dumps(incomplete_payload), encoding="utf-8")

        # When / Then
        with self.assertRaises(IncompleteRunRecordError):
            build_results_workbook(records, self.root / "out.xlsx", REFERENCE_WORKBOOK)

    def test_refuses_to_overwrite_template(self) -> None:
        # Given
        records = self._records()
        copied_template = self.root / "template.xlsx"
        shutil.copy2(REFERENCE_WORKBOOK, copied_template)
        template_hash = hashlib.sha256(copied_template.read_bytes()).hexdigest()

        # When / Then
        with self.assertRaises(RerunRecordError):
            build_results_workbook(records, copied_template, copied_template)
        self.assertEqual(template_hash, hashlib.sha256(copied_template.read_bytes()).hexdigest())

    def _records(self, failed_source_id: Optional[str] = None) -> Tuple[SourceRecord, ...]:
        records: List[SourceRecord] = []
        for source_index, source_id in enumerate(SOURCE_IDS):
            run_directory = self.root / source_id
            config_path = run_directory / "config.json"
            results_path = run_directory / "test_results.json"
            run_directory.mkdir()
            config_path.write_text("{}", encoding="utf-8")
            if source_id == failed_source_id:
                records.append(
                    SourceRecord(
                        source_id=source_id,
                        status="failed",
                        run_directory=run_directory,
                        config_path=config_path,
                        test_results_path=results_path,
                        reason="GPU memory exhausted",
                    )
                )
                continue
            payload = self._payload(source_index)
            self.payloads[source_id] = payload
            results_path.write_text(json.dumps(payload), encoding="utf-8")
            records.append(
                SourceRecord(
                    source_id=source_id,
                    status="success",
                    run_directory=run_directory,
                    config_path=config_path,
                    test_results_path=results_path,
                    reason=None,
                )
            )
        return tuple(records)

    def _payload(self, source_index: int) -> Dict[str, Dict[str, Union[List[int], List[List[float]]]]]:
        labels = [0, 1, 2, 0, 1, 2]
        prediction_sets = (
            [0, 1, 2, 0, 1, 2],
            [0, 2, 1, 0, 2, 1],
            [1, 1, 2, 0, 0, 2],
            [2, 1, 0, 2, 1, 0],
        )
        payload: Dict[str, Dict[str, Union[List[int], List[List[float]]]]] = {}
        for split_index, split_name in enumerate(SPLITS):
            predictions = prediction_sets[(source_index + split_index) % len(prediction_sets)]
            probabilities = [[0.8 if label == prediction else 0.1 for label in range(3)] for prediction in predictions]
            payload[split_name] = {
                "labels": labels,
                "predictions": predictions,
                "probabilities": probabilities,
            }
        return payload

    def _assert_summary_values(self, workbook) -> None:
        base_bilstm = self.payloads["BASE_BILSTM"]
        full_bilstm = self.payloads["FULL_BILSTM"]
        task_bilstm = self.payloads["TASK_BILSTM"]
        base_metrics = {split: compute_metrics(base_bilstm[split], split) for split in SPLITS}
        full_metrics = {split: compute_metrics(full_bilstm[split], split) for split in SPLITS}
        task_metrics = {split: compute_metrics(task_bilstm[split], split) for split in SPLITS}
        base_ood_f1 = sum(base_metrics[split]["F1 Macro"] for split in SPLITS[1:]) / 3
        full_ood_f1 = sum(full_metrics[split]["F1 Macro"] for split in SPLITS[1:]) / 3
        task_ood_acc = sum(task_metrics[split]["ACC"] for split in SPLITS[1:]) / 3
        distribution = workbook["Distribution Shift"]
        configurations = workbook["BiLSTM Configurations"]
        components = workbook["Component Ablation"]
        encoders = workbook["Encoder Comparison"]
        self.assertAlmostEqual(distribution["B5"].value, base_metrics["Val"]["ACC"])
        self.assertAlmostEqual(configurations["F5"].value, base_ood_f1)
        self.assertAlmostEqual(components["C6"].value, task_ood_acc)
        self.assertAlmostEqual(encoders["B10"].value, base_ood_f1)
        self.assertAlmostEqual(encoders["C10"].value, full_ood_f1)
        self.assertAlmostEqual(encoders["D10"].value, (full_ood_f1 - base_ood_f1) * 100)
        self.assertEqual([encoders.cell(6, column).value for column in range(2, 8)], [None] * 6)


if __name__ == "__main__":
    unittest.main()
