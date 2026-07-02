# -*- coding: utf-8 -*-
"""Tests for Schulte raw validation and inference plumbing."""

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.inference.schulte_raw_loader import build_schulte_records
from src.inference.schulte_raw_question import load_schulte_question_info
from src.inference.schulte_raw_runner import (
    run_schulte_raw_inference,
    run_schulte_raw_validate,
)


class SchulteRawRunnerTest(unittest.TestCase):
    def test_question_info_maps_training_task_conditions(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            question_path = Path(temp_dir) / "题目信息.xlsx"
            write_question_info(question_path)

            questions = load_schulte_question_info(question_path)

        self.assertEqual(questions[1].task_conditions, (1, 1, 0, 0, 0))
        self.assertEqual(questions[2].task_conditions, (1, 0, 1, 0, 0))
        self.assertEqual(questions[29].task_conditions, (4, 1, 1, 1, 1))
        self.assertEqual(questions[30].task_conditions, (4, 1, 0, 1, 1))

    def test_build_records_reads_raw_xlsx_without_labels(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir, question_path = build_raw_dataset(root, task_ids=(1,))

            build = build_schulte_records(data_dir, question_path, task_ids=(1,))
            record = build.records[0]

        self.assertFalse(build.issues)
        self.assertEqual(record.status, "ready")
        self.assertEqual(record.subject_id, "0429001")
        self.assertEqual(record.task_id, 1)
        self.assertEqual(record.click_count, 3)
        self.assertEqual(record.segment_count, 2)
        self.assertEqual(record.sample.task_conditions, (1, 1, 0, 0, 0))
        self.assertFalse(hasattr(record.sample, "label"))

    def test_timestamp_fix_and_coordinate_clipping_are_reported(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir, question_path = build_raw_dataset(
                root,
                task_ids=(1,),
                rollback=True,
                out_of_bounds=True,
            )

            build = build_schulte_records(data_dir, question_path, task_ids=(1,))
            record = build.records[0]

        self.assertEqual(record.timestamp_fix_count, 1)
        self.assertEqual(record.clipped_coordinate_count, 1)
        self.assertTrue(any("时间回退" in warning for warning in record.warnings))
        self.assertTrue(any("坐标裁剪" in warning for warning in record.warnings))

    def test_unexpected_time_rollback_becomes_validate_issue(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir, question_path = build_raw_dataset(
                root,
                task_ids=(1,),
                bad_rollback=True,
            )

            build = build_schulte_records(data_dir, question_path, task_ids=(1,))

        self.assertEqual(build.records[0].status, "invalid")
        self.assertEqual(len(build.issues), 1)
        self.assertIn("无法自动修正", build.issues[0].reason)

    def test_validate_writes_csv_xlsx_and_issues_sheet(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir, question_path = build_raw_dataset(root, task_ids=(1,))

            output = run_schulte_raw_validate(
                data_dir,
                question_path,
                output_root=root / "outputs",
                task_ids=(1,),
            )
            csv_path = output / "manifest.csv"
            csv_exists = csv_path.is_file()
            header = read_csv_header(csv_path)
            workbook = load_workbook(output / "manifest.xlsx", read_only=True)
            sheetnames = tuple(workbook.sheetnames)
            workbook.close()

        self.assertTrue(csv_exists)
        self.assertIn("timestamp_fix_count", header)
        self.assertIn("summary", sheetnames)
        self.assertIn("manifest", sheetnames)
        self.assertIn("issues", sheetnames)
        self.assertIn("q01", sheetnames)

    def test_infer_requires_explicit_checkpoint(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir, question_path = build_raw_dataset(root, task_ids=(1,))

            with self.assertRaisesRegex(ValueError, "checkpoint"):
                run_schulte_raw_inference(
                    data_dir,
                    question_path,
                    checkpoint_path=None,
                    output_root=root / "outputs",
                    task_ids=(1,),
                )


def build_raw_dataset(root: Path, task_ids, **options):
    data_dir = root / "舒尔特-原始"
    subject_dir = data_dir / "0429001"
    subject_dir.mkdir(parents=True)
    question_path = root / "题目信息.xlsx"
    write_question_info(question_path)
    for task_id in task_ids:
        write_task_xlsx(subject_dir / f"{task_id}.xlsx", task_id, **options)
    return data_dir, question_path


def write_question_info(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(QUESTION_HEADERS)
    for task_id in range(1, 31):
        sheet.append(question_row(task_id))
    workbook.save(path)


def question_row(task_id: int):
    grid_text = grid_text_for_task(task_id)
    number_range = "1-99" if task_id in (1, 3, 5, 29, 30) else "1-N"
    click_disappear = "✔" if task_id in (2, 3, 6, 27, 29) else "×"
    grid_distractors = 3 if task_id in (5, 6, 29, 30) else 0
    number_distractors = 9 if task_id in (3, 4, 5, 6, 29, 30) else 0
    has_distractor = "是" if grid_distractors or number_distractors else "否"
    return [
        task_id,
        grid_text,
        number_range,
        click_disappear,
        has_distractor,
        grid_distractors,
        number_distractors,
        grid_distractors + number_distractors,
    ]


def grid_text_for_task(task_id: int) -> str:
    if task_id <= 6:
        return "3×3"
    if task_id <= 14:
        return "4×4"
    if task_id <= 26:
        return "5×5"
    return "6×6"


def write_task_xlsx(
    path: Path,
    task_id: int,
    *,
    rollback: bool = False,
    bad_rollback: bool = False,
    out_of_bounds: bool = False,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_level_sheet(workbook, task_id)
    write_grid_sheet(workbook)
    write_operation_sheet(workbook)
    write_gaze_sheet(
        workbook,
        rollback=rollback,
        bad_rollback=bad_rollback,
        out_of_bounds=out_of_bounds,
    )
    workbook.save(path)


def write_level_sheet(workbook: Workbook, task_id: int) -> None:
    sheet = workbook.create_sheet("关卡信息")
    sheet.append(["关卡名称", "网格数量", "附加难度", "耗时", "错误次数"])
    sheet.append([f"题目{task_id}", 3, "None", 3.0, 1])


def write_grid_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("网格信息")
    sheet.append(["网格序号", "数字", "坐标X", "坐标Y", "边长"])
    for number in range(1, 10):
        sheet.append([number, number, number * 10, number * 10, 20])


def write_operation_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("操作")
    sheet.append(["时间", "正确", "坐标X", "坐标Y"])
    sheet.append(["2026/5/10 14:21:48:000", "正确", 100, 100])
    sheet.append(["2026/5/10 14:21:49:000", "错误", 200, 200])
    sheet.append(["2026/5/10 14:21:50:000", "正确", 300, 300])


def write_gaze_sheet(
    workbook: Workbook,
    *,
    rollback: bool,
    bad_rollback: bool,
    out_of_bounds: bool,
) -> None:
    sheet = workbook.create_sheet("视线焦点")
    sheet.append(["时间", "坐标X", "坐标Y"])
    rows = gaze_rows(rollback, bad_rollback, out_of_bounds)
    for row in rows:
        sheet.append(row)


def gaze_rows(rollback: bool, bad_rollback: bool, out_of_bounds: bool):
    if bad_rollback:
        return [
            ["2026/5/10 14:21:48:964", 100, 100],
            ["2026/5/10 14:21:47:000", 110, 110],
        ]
    if rollback:
        x_value = 2000 if out_of_bounds else 110
        return [
            ["2026/5/10 14:21:48:964", 100, 100],
            ["2026/5/10 14:21:48:000", x_value, -5 if out_of_bounds else 110],
            ["2026/5/10 14:21:49:018", 120, 120],
            ["2026/5/10 14:21:49:500", 130, 130],
        ]
    x_value = 2000 if out_of_bounds else 110
    return [
        ["2026/5/10 14:21:48:100", 100, 100],
        ["2026/5/10 14:21:48:500", x_value, -5 if out_of_bounds else 110],
        ["2026/5/10 14:21:49:500", 120, 120],
        ["2026/5/10 14:21:50:100", 130, 130],
    ]


def read_csv_header(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return next(csv.reader(file_obj))


QUESTION_HEADERS = [
    "题目",
    "方格数量",
    "数字范围",
    "点击是否消失",
    "是否有干扰项",
    "方格干扰项数量",
    "数字干扰项数量",
    "干扰项总数量",
]


if __name__ == "__main__":
    unittest.main()
